# Erzeugt test/fixtures/Beispielreise2030.xlsx – eine erfundene Ferien-Excel im Aufbau der früheren Vorlage.
# Nur für die automatischen Tests des Excel-Imports. Aufruf: python3 test/fixtures/make-beispiel-excel.py
import datetime as dt
from pathlib import Path
from openpyxl import Workbook

wb = Workbook()
ue = wb.active
ue.title = 'Übersicht'
ue.append(['Tag', 'Datum', 'Programm', 'Ort', 'Übernachtung', 'Notiz'])
start = dt.date(2030, 3, 1)
plan = [
    ('Anreise', 'Lissabon', 'Hotel Sol', 'Flug am Morgen'),
    ('Altstadt', 'Lissabon', 'Hotel Sol', ''),
    ('Fahrt an die Küste', 'Lagos', 'Casa Mar', ''),
    ('Tauchen', 'Lagos', 'Casa Mar', 'Zwei Tauchgänge'),
    ('Rückreise', '', '', ''),
]
for i, (prog, ort, stay, note) in enumerate(plan):
    ue.append([i + 1, start + dt.timedelta(days=i), prog, ort, stay, note])

ak = wb.create_sheet('Allgemeine Kosteneinstellung')
ak.append(['Pauschale', 'Betrag CHF'])
ak.append(['Essen & Trinken', 70])
ak.append(['Unterkunft', 120])

wr = wb.create_sheet('Währungsrechner')
wr.append(['Währung', 'CHF pro Einheit'])
wr.append(['EUR', 0.95])

ra = wb.create_sheet('Reiseablauf')
ra.append(['Datum', 'Aktivität', 'Unterkunft', 'Betrag', 'Währung', 'CHF', 'Bezahlmethode', 'Status', 'Kommentar'])
d = lambda k: start + dt.timedelta(days=k)
rows = [
    [d(0), 'Flug Zürich–Lissabon', '07:15', 420, 'CHF', 420, 'Testkarte', 'bezahlt', ''],
    [None, 'Unterkunft', 'Hotel Sol', 180, 'EUR', 171, 'Testkarte', 'reserviert', 'zahlbar bis 15.2.30'],
    [None, 'Essen & Trinken', '', None, 'CHF', None, '', '', ''],
    [d(1), 'Unterkunft', 'Hotel Sol', 0, 'EUR', 0, '', 'reserviert', ''],
    [d(1), 'Stadtführung', '10:00', 35, 'EUR', None, '', 'vor Ort', 'https://example.org/tour'],
    [d(2), 'Mietauto', '', 150, 'EUR', None, 'Testkarte', 'angezahlt', ''],
    [d(2), 'Unterkunft', 'Casa Mar', 200, 'EUR', None, '', 'vor Ort', ''],
    [d(3), 'Unterkunft', 'Casa Mar', 0, 'EUR', None, '', 'vor Ort', ''],
    [d(3), 'Tauchen 2 Tauchgänge', '', 110, 'EUR', None, '', 'vor Ort', ''],
    [d(4), 'Flug Lissabon–Zürich', '18:40', 0, 'CHF', None, '', '', 'im Hinflug enthalten'],
    [d(9), 'Ausserhalb', '', 10, 'CHF', None, '', '', ''],
]
for r in rows:
    ra.append(r)
# Pauschale per Formel auf die Kosteneinstellung (wie in der Vorlage)
ra['D4'] = "='Allgemeine Kosteneinstellung'!$B$2"

ab = wb.create_sheet('Abrechnung')
ab.append(['Was', 'Betrag'])
ab.append(['Testkarte', '=SUMIF(Reiseablauf!G:G,A2,Reiseablauf!F:F)'])
ab.append(['Total', 591])
ab.append(['pro Person', '=B3/2'])
ab.append(['bereits bezahlt', 420])

out = Path(__file__).with_name('Beispielreise2030.xlsx')
wb.save(out)
print('geschrieben:', out)
